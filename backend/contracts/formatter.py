import os
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(16)


def _add_label_value(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    paragraph.add_run(value)


def _add_signature_block(doc: Document, party: str) -> None:
    doc.add_paragraph()
    doc.add_paragraph("______________________________")
    doc.add_paragraph(f"Signature ({party})")
    doc.add_paragraph("Name:")
    doc.add_paragraph("Title:")
    doc.add_paragraph("Date:")


def build_contract_doc(
    title: str,
    body: str,
    party_a: Optional[str] = None,
    party_b: Optional[str] = None,
    effective_date: Optional[str] = None,
) -> Document:
    doc = Document()
    _add_heading(doc, title)
    doc.add_paragraph()
    _add_label_value(
        doc,
        "Effective date",
        effective_date or datetime.utcnow().strftime("%Y-%m-%d"),
    )
    if party_a:
        _add_label_value(doc, "Party A", party_a)
    if party_b:
        _add_label_value(doc, "Party B", party_b)

    doc.add_paragraph()
    for line in body.splitlines():
        doc.add_paragraph(line)

    doc.add_paragraph()
    doc.add_paragraph("Signatures")
    _add_signature_block(doc, party_a or "Party A")
    _add_signature_block(doc, party_b or "Party B")
    return doc


def save_contract_doc(doc: Document, filename: str, storage_path: str) -> str:
    os.makedirs(storage_path, exist_ok=True)
    output_path = os.path.join(storage_path, filename)
    doc.save(output_path)
    return output_path


def save_contract_pdf(
    title: str,
    body: str,
    storage_path: str,
    filename: str,
    party_a: Optional[str] = None,
    party_b: Optional[str] = None,
    effective_date: Optional[str] = None,
) -> str:
    os.makedirs(storage_path, exist_ok=True)
    output_path = os.path.join(storage_path, filename)
    styles = getSampleStyleSheet()
    title_style = styles["Heading1"]
    body_style = styles["BodyText"]

    doc = SimpleDocTemplate(output_path, pagesize=LETTER)
    content = []
    content.append(Paragraph(title, title_style))
    content.append(Spacer(1, 12))
    content.append(
        Paragraph(
            f"Effective date: {effective_date or datetime.utcnow().strftime('%Y-%m-%d')}",
            body_style,
        )
    )
    if party_a:
        content.append(Paragraph(f"Party A: {party_a}", body_style))
    if party_b:
        content.append(Paragraph(f"Party B: {party_b}", body_style))
    content.append(Spacer(1, 12))

    for line in body.splitlines():
        if line.strip():
            content.append(Paragraph(line, body_style))
            content.append(Spacer(1, 6))

    content.append(Spacer(1, 12))
    content.append(Paragraph("Signatures", body_style))
    content.append(Spacer(1, 12))
    for party in [party_a or "Party A", party_b or "Party B"]:
        content.append(Paragraph("______________________________", body_style))
        content.append(Paragraph(f"Signature ({party})", body_style))
        content.append(Paragraph("Name:", body_style))
        content.append(Paragraph("Title:", body_style))
        content.append(Paragraph("Date:", body_style))
        content.append(Spacer(1, 12))

    doc.build(content)
    return output_path
